import hashlib
import os
import time
import math, random 
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill, Font, Border, Side, Alignment
from django.conf import settings

def generate_class_code(total_digits,existing_codes) :  
    digits = ''.join([str(i) for i in range(0,10)])
    code = ""  
    while True:
        for i in range(total_digits) : 
            code += digits[math.floor(random.random() * 10)] 
        if code not in existing_codes:
            # print('Code not in existing codes')
            break
    return code 

def generate_unique_id(phone_number, email):
    # Concatenate phone number and email
    unique_string = f"{email}{phone_number}"
    
    # Generate a salt
    salt = os.urandom(16).hex()
    
    # Get the current timestamp
    timestamp = str(int(time.time()))
    
    # Combine the unique string, salt, and timestamp
    combined_string = f"{unique_string}{salt}{timestamp}"
    
    # Use hashlib to generate a unique hash
    hashed_string = hashlib.sha512(combined_string.encode()).hexdigest()
    
    # Return the first 10 characters of the hash with a prefix
    return 'IN' + hashed_string[:10].upper()

def generate_unique_employeeid(phone_number, email):
    # Concatenate phone number and email
    unique_string = f"{email}{phone_number}"
    
    # Generate a salt
    salt = os.urandom(16).hex()
    
    # Get the current timestamp
    timestamp = str(int(time.time()))
    
    # Combine the unique string, salt, and timestamp
    combined_string = f"{unique_string}{salt}{timestamp}"
    
    # Use hashlib to generate a unique hash
    hashed_string = hashlib.sha512(combined_string.encode()).hexdigest()
    
    # Return the first 10 characters of the hash with a prefix
    return 'EMP' + hashed_string[:10].upper()


def generate_unique_establishment(phone_number, email):
    # Concatenate phone number and email
    unique_string = f"{email}{phone_number}"
    
    # Generate a salt
    salt = os.urandom(16).hex()
    
    # Get the current timestamp
    timestamp = str(int(time.time()))
    
    # Combine the unique string, salt, and timestamp
    combined_string = f"{unique_string}{salt}{timestamp}"
    
    # Use hashlib to generate a unique hash
    hashed_string = hashlib.sha512(combined_string.encode()).hexdigest()
    
    # Return the first 10 characters of the hash with a prefix
    return 'ES' + hashed_string[:10].upper()

def generate_unique_ngo(phone_number, email):
    # Concatenate phone number and email
    unique_string = f"{email}{phone_number}"
    
    # Generate a salt
    salt = os.urandom(16).hex()
    
    # Get the current timestamp
    timestamp = str(int(time.time()))
    
    # Combine the unique string, salt, and timestamp
    combined_string = f"{unique_string}{salt}{timestamp}"
    
    # Use hashlib to generate a unique hash
    hashed_string = hashlib.sha512(combined_string.encode()).hexdigest()
    
    # Return the first 10 characters of the hash with a prefix
    return 'NGO' + hashed_string[:10].upper()

def generate_unique_consultancy(phone_number, email):
    # Concatenate phone number and email
    unique_string = f"{email}{phone_number}"
    
    # Generate a salt
    salt = os.urandom(16).hex()
    
    # Get the current timestamp
    timestamp = str(int(time.time()))
    
    # Combine the unique string, salt, and timestamp
    combined_string = f"{unique_string}{salt}{timestamp}"
    
    # Use hashlib to generate a unique hash
    hashed_string = hashlib.sha512(combined_string.encode()).hexdigest()
    
    # Return the first 10 characters of the hash with a prefix
    return 'CON' + hashed_string[:10].upper()


# Utility function to update session data
def update_locations_session(request, locations_data):

    if 'locations_data' in request.session:
        del request.session['locations_data']
    
    request.session['locations_data'] = locations_data
    # print(request.session.get('locations_data'))

def get_session_data(request):
    return request.session.get('locations_data')

# Utility function to update session data
def update_vendor_session(request, vendor_data):

    if 'vendor_data' in request.session:
        del request.session['vendor_data']
    
    request.session['vendor_data'] = vendor_data
    # print(request.session.get('locations_data'))

def get_vendor_data(request):
    return request.session.get('vendor_data')


def create_vendor_excel(locations):

    # Colors and styles for cell formatting
    brown_fill = PatternFill(start_color="FABF8F", end_color="FABF8F", fill_type="solid")
    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

    # Border style
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    # Alignment style
    alignment_left_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)
    alignment_center_wrap = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Create workbook and worksheet
    wb = Workbook()
    ws = wb.active

    # Set column widths for better readability
    column_widths = [12, 35, 20, 30, 22, 25, 20, 25, 18, 25, 20, 20, 30]
    for i, width in enumerate(column_widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = width 

    headers = [
        "SR.NO", "VENDOR NAME", "VENDOR'S MYPOSH UID", "VENDOR'S COMMUNICATION ADDRESS*", "MOBILE NUMBER*", "EMAIL ID*", "NATURE OF SERVICE",
        "CONTACT PERSON NAME*", "CONTACT PERSON MOBILE NO*", "CONTACT PERSON EMAIL ID*", "CONTRACT COMMENCEMENT DATE",
        "CONTRACT EXPIRY DATE", "MAX NUMBER OF EMPLOYEES DEPLOYED"
    ]

    row_num = 1
    for location in locations:
        # Set "LOCATION" label and location name with border and alignment
        location_label_cell = ws.cell(row=row_num, column=1, value="LOCATION")
        location_name_cell = ws.cell(row=row_num, column=2, value=location["location_name"])

        for i in range(3, 14):
            location_name_color = ws.cell(row=row_num, column=i, value="")
            location_name_color.fill = brown_fill
            location_name_color.border = thin_border

        location_label_cell.fill = brown_fill
        location_name_cell.fill = brown_fill
        location_label_cell.border = thin_border
        location_name_cell.border = thin_border
        location_label_cell.alignment = alignment_left_wrap
        location_name_cell.alignment = alignment_left_wrap
        
        # Move to the next row to add headers
        row_num += 1
        
        # Fill header row with brown color, border, and alignment
        for col_num, header in enumerate(headers, start=1):
            cell = ws.cell(row=row_num, column=col_num, value=header)
            cell.fill = brown_fill
            cell.border = thin_border
            cell.alignment = alignment_center_wrap

        # Move to the next row for vendor entries
        row_num += 1
        
        # Add vendors under the location
        for sr_no in range(1, location["vendors"] + 1):
            ws.cell(row=row_num, column=1, value=sr_no).fill = brown_fill
            
            # Fill all vendor cells with yellow color, border, and alignment
            for col_num in range(1, 14):
                cell = ws.cell(row=row_num, column=col_num)
                cell.fill = yellow_fill
                cell.border = thin_border
                cell.alignment = alignment_left_wrap
                if col_num == 1:
                    cell.value = sr_no  # Set the serial number
                    cell.alignment = alignment_center_wrap
                if col_num == 5 or col_num == 9 or col_num == 13:
                    cell.number_format = '0'
                elif col_num == 11 or col_num == 12:
                    cell.number_format = 'yyyy-mm-dd'
        
            # Move to the next row for the next vendor
            row_num += 1
        
        row_num += 1

    file_path = os.path.join(settings.BASE_DIR, 'static/posh/')
    
    wb.save(f"{file_path}VendorDataTemplate.xlsx")


def create_employee_table(data):

    brown_fill = PatternFill(start_color="FABF8F", end_color="FABF8F", fill_type="solid")

    # Border style
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    # Alignment style
    alignment_center_wrap = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Load your workbook and select the sheet
    file_path = os.path.join(settings.BASE_DIR, 'static/posh/Employee Templete.xlsx')  # Replace with the actual file path
    workbook = load_workbook(filename=file_path)
    sheet = workbook.active  # Select the active sheet, or use workbook['SheetName'] if you know the sheet name

    # Starting row for data entry (assuming headers are on the first row)
    start_row = 2

    # Iterate over each location in the data dictionary
    for location_name, details in data.items():
        # Fill rows for Direct Employees
        direct_employee_count = details.get("Direct Employee", 0)
        for _ in range(direct_employee_count):
            # SR.NO
            cell_sr_no = sheet.cell(row=start_row, column=1, value=start_row - 1)
            # Location Name
            cell_location = sheet.cell(row=start_row, column=2, value=location_name)
            # Nature of Employment
            cell_employment = sheet.cell(row=start_row, column=3, value="Direct")
            # Vendor (empty for direct employees)
            cell_vendor = sheet.cell(row=start_row, column=4, value="")

            # Apply styles to each cell
            for cell in [cell_sr_no, cell_location, cell_employment, cell_vendor]:
                cell.fill = brown_fill
                cell.border = thin_border
                cell.alignment = alignment_center_wrap

            start_row += 1

        # Fill rows for each Vendor
        for vendor in details.get("Vendors", []):
            vendor_name = vendor.get("Vendor Name", "")
            max_employees = vendor.get("Max Employees", 0)

            # Repeat for the number of Max Employees for each vendor
            for _ in range(max_employees):
                # SR.NO
                cell_sr_no = sheet.cell(row=start_row, column=1, value=start_row - 1)
                # Location Name
                cell_location = sheet.cell(row=start_row, column=2, value=location_name)
                # Nature of Employment
                cell_employment = sheet.cell(row=start_row, column=3, value="Indirect")
                # Vendor Name
                cell_vendor = sheet.cell(row=start_row, column=4, value=vendor_name)

                # Apply styles to each cell
                for cell in [cell_sr_no, cell_location, cell_employment, cell_vendor]:
                    cell.fill = brown_fill
                    cell.border = thin_border
                    cell.alignment = alignment_center_wrap

                start_row += 1

    # Save the workbook after filling data
    file_path = os.path.join(settings.BASE_DIR, 'static/posh/')
    
    workbook.save(f"{file_path}EmployeeDataTemplate.xlsx")
